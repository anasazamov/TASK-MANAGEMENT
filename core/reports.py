"""Employee task statistics, shown on the employees page and exported to Excel."""
import io
from datetime import date, datetime, time, timedelta

from django.db.models import Count, Q
from django.utils import timezone

from .models import Task
from .permissions import assignees_for

COLUMNS = [('T/R', 7), ('F.I.Sh.', 50), ('Topshiriqlar soni', 15), ('Bajarilmoqda', 17),
           ('Bajarilmagan', 19), ('Tanishilmagan topshiriqlar', 19)]


def parse_week(value):
    """'2026-W38' (HTML week input) -> (monday, next monday) as aware datetimes, else None."""
    try:
        year, week = value.split('-W')
        monday = date.fromisocalendar(int(year), int(week), 1)
    except (AttributeError, ValueError):
        return None
    start = timezone.make_aware(datetime.combine(monday, time.min))
    return start, start + timedelta(days=7)


def current_week():
    year, week, _ = timezone.localdate().isocalendar()
    return f'{year}-W{week:02d}'


def employee_stats(user, period=None):
    tasks = Task.objects.visible_to(user)
    if period:
        tasks = tasks.filter(created_at__gte=period[0], created_at__lt=period[1])
    now = timezone.now()
    overdue = Q(status='active', due_at__lt=now)
    rows = list(tasks.values('assignee_id', 'assignee__full_name').annotate(
        total=Count('pk'),
        in_progress=Count('pk', filter=Q(status='submitted') | (Q(status='active') & ~overdue)),
        overdue=Count('pk', filter=overdue),
        unseen=Count('pk', filter=Q(seen_at__isnull=True)),
    ))
    # People this user may assign belong in the workload even with nothing on them.
    counted = {row['assignee_id'] for row in rows}
    rows += [{'assignee_id': person.pk, 'assignee__full_name': person.full_name,
              'total': 0, 'in_progress': 0, 'overdue': 0, 'unseen': 0}
             for person in assignees_for(user).exclude(pk__in=counted)]
    return sorted(rows, key=lambda row: (-row['total'], row['assignee__full_name']))


def period_label(period):
    if not period:
        return 'Barcha vaqt'
    last = timezone.localtime(period[1]) - timedelta(days=1)
    return f"{timezone.localtime(period[0]):%d.%m.%Y}–{last:%d.%m.%Y}"


def workbook(rows, period, prepared_by):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = 'Xodimlar statistikasi'
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    colors = {4: '00B050', 5: 'FF0000'}
    for index, (_, width) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet['C1'] = 'Xodimlar statistikasi'
    sheet['C1'].font = Font(name='Times New Roman', size=14, bold=True)
    sheet['C2'] = f'Davr: {period_label(period)}'
    sheet['C3'] = f'{timezone.localdate():%d.%m.%Y}-yil'
    for cell in ('C2', 'C3'):
        sheet[cell].font = Font(name='Times New Roman', size=14)
    for index, (title, _) in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=5, column=index, value=title)
        cell.font = Font(name='Times New Roman', size=12, bold=True, color=colors.get(index))
        cell.alignment, cell.border = center, border
    for number, row in enumerate(rows, 1):
        values = [number, row['assignee__full_name'], row['total'], row['in_progress'], row['overdue'], row['unseen']]
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row=5+number, column=index, value=value)
            cell.font = Font(name='Times New Roman', size=14, color=colors.get(index))
            cell.alignment = Alignment(horizontal='left' if index == 2 else 'center', vertical='center', wrap_text=index == 2)
            cell.border = border
    sheet.cell(row=5+len(rows)+3, column=3, value=prepared_by).font = Font(name='Times New Roman', size=14)
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


ATTENDANCE_STATUS = {'present': 'Keldi', 'review': 'Tekshiruvda', 'rejected': 'Rad etildi'}
ATTENDANCE_COLUMNS = [('T/R', 7), ('F.I.Sh.', 40), ('Bo‘linma', 26), ('Kirishlar soni', 15),
                      ('Birinchi kirish', 16), ('Oxirgi kirish', 16), ('Kechikish', 16), ('Holat', 16),
                      ('Qaror / sabab', 44), ('Uzrli sabab', 36)]


# Row colours for the marked export: light fills that still read when printed.
ATTENDANCE_FILLS = {
    'absent': ('FFC7CE', 'Kelmadi'),
    'late': ('FFEB9C', 'Kechikdi'),
    'rejected': ('F4B183', 'Rad etildi'),
    'review': ('FFF2CC', 'Tekshiruvda'),
    'excused': ('DDEBF7', 'Uzrli'),
}


def attendance_state(row):
    """One word for the row: present, late, review, rejected, excused or absent."""
    attendance = row['attendance']
    if not attendance:
        return 'excused' if row.get('excuse') else 'absent'
    if attendance.status in ('review', 'rejected'):
        return attendance.status
    return 'late' if row.get('late') else 'present'


def attendance_workbook(rows, day, prepared_by, marked=False):
    """Daily attendance (one row per employee) as an .xlsx, like the stats sheet.

    `marked` colours each row by its state — absentees red — and adds a legend
    with the counts, so the sheet can be read at a glance or printed as is.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = 'Davomat'
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    fills = {state: PatternFill('solid', start_color=color, end_color=color)
             for state, (color, _) in ATTENDANCE_FILLS.items()}
    for index, (_, width) in enumerate(ATTENDANCE_COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet['C1'] = 'Kunlik davomat'
    sheet['C1'].font = Font(name='Times New Roman', size=14, bold=True)
    sheet['C2'] = f'Sana: {day:%d.%m.%Y}'
    sheet['C2'].font = Font(name='Times New Roman', size=14)
    for index, (title, _) in enumerate(ATTENDANCE_COLUMNS, 1):
        cell = sheet.cell(row=4, column=index, value=title)
        cell.font = Font(name='Times New Roman', size=12, bold=True)
        cell.alignment, cell.border = center, border
    counts = {}
    for number, row in enumerate(rows, 1):
        attendance = row['attendance']
        state = attendance_state(row)
        counts[state] = counts.get(state, 0) + 1
        values = [number, row['user'].full_name, str(row['user'].department or '—'),
                  attendance.entry_count if attendance else 0,
                  timezone.localtime(attendance.first_entry_at).strftime('%H:%M') if attendance and attendance.first_entry_at else '—',
                  timezone.localtime(attendance.last_entry_at).strftime('%H:%M') if attendance and attendance.last_entry_at else '—',
                  f"+{row['late_label']}" if row.get('late') else '—']
        excuse = row.get('excuse')
        if not attendance:
            values.append('Uzrli' if excuse else 'Kelmadi')
        else:
            values.append('Kechikdi' if row.get('late') else ATTENDANCE_STATUS.get(attendance.status, ''))
        review = row.get('review')
        values.append(f'{review.get_state_display()} ({review.resolved_by.short_name if review.resolved_by else "avtomatik"}, '
                      f'{timezone.localtime(review.resolved_at):%d.%m.%Y %H:%M})'
                      + (f': {review.reason}' if review.reason else '') if review else '')
        values.append(excuse.label if excuse else '')
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row=4 + number, column=index, value=value)
            cell.font = Font(name='Times New Roman', size=13, bold=marked and state == 'absent' and index == 8)
            text = index in (2, 9, 10)
            cell.alignment = Alignment(horizontal='left' if text else 'center', vertical='center', wrap_text=text)
            cell.border = border
            if marked and state in fills:
                cell.fill = fills[state]
    last = 4 + len(rows)
    if marked:
        # Legend: colour, meaning and how many, under the table.
        sheet.cell(row=last + 2, column=2, value='Belgilar:').font = Font(name='Times New Roman', size=12, bold=True)
        legend = [('present', None, 'Keldi')] + [(state, fills[state], label)
                                                 for state, (_, label) in ATTENDANCE_FILLS.items()]
        for offset, (state, fill, label) in enumerate(legend):
            swatch = sheet.cell(row=last + 3 + offset, column=2, value=f'{label}: {counts.get(state, 0)}')
            swatch.font = Font(name='Times New Roman', size=12)
            swatch.border = border
            if fill:
                swatch.fill = fill
        last += len(legend) + 2
    sheet.cell(row=last + 3, column=3, value=prepared_by).font = Font(name='Times New Roman', size=13)
    sheet.print_options.horizontalCentered = True
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
