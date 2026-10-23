import io
import json
import tempfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone

from .attendance import ingest_event, resolve_review
from .models import (AccessCard, AccessEvent, Attendance, AttendanceAlert, AttendanceExcuse,
                     FaceEnrollment, FaceReview, User)


def event(time, card='1000000001', employee='13', name='SINOV XODIM', minor=1, major=5, door=1, serial=None):
    data = {'major': major, 'minor': minor, 'time': time, 'cardNo': card,
            'employeeNoString': employee, 'name': name, 'doorNo': door}
    if serial is not None:
        data['serialNo'] = serial
    return data


@override_settings(ATTENDANCE_DAY_START_HOUR=7, ATTENDANCE_DOOR_MAJOR=5)
class IngestTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')
        self.card = AccessCard.objects.create(user=self.user, card_no='1000000001', device_employee_no='13')

    def test_card_event_creates_event_and_daily_attendance(self):
        result = ingest_event(event('2026-10-21 08:15:30'))
        self.assertTrue(result['stored'])
        self.assertTrue(result['matched'])
        self.assertTrue(result['first_entry'])
        access = AccessEvent.objects.get(pk=result['event_id'])
        self.assertEqual(access.user, self.user)
        day = Attendance.objects.get(user=self.user, date=access.happened_at.astimezone(timezone.get_current_timezone()).date())
        self.assertEqual(day.entry_count, 1)
        self.assertEqual(day.first_entry_at, access.happened_at)

    def test_entries_accumulate_first_and_last(self):
        ingest_event(event('2026-10-21 08:00:00'))
        ingest_event(event('2026-10-21 12:30:00'))
        last = ingest_event(event('2026-10-21 17:45:00'))
        self.assertFalse(last['first_entry'])
        day = Attendance.objects.get(user=self.user, date=timezone.localtime(AccessEvent.objects.earliest('happened_at').happened_at).date())
        self.assertEqual(day.entry_count, 3)
        self.assertEqual(timezone.localtime(day.first_entry_at).strftime('%H:%M'), '08:00')
        self.assertEqual(timezone.localtime(day.last_entry_at).strftime('%H:%M'), '17:45')

    def test_same_event_is_not_counted_twice(self):
        one = event('2026-10-21 08:00:00', serial='555')
        self.assertTrue(ingest_event(one)['stored'])
        self.assertFalse(ingest_event(one)['stored'])
        self.assertEqual(AccessEvent.objects.count(), 1)
        self.assertEqual(Attendance.objects.get(user=self.user).entry_count, 1)

    def test_entry_before_day_start_is_logged_but_not_counted(self):
        result = ingest_event(event('2026-10-21 06:30:00'))
        self.assertTrue(result['stored'])
        self.assertFalse(result.get('counted'))
        self.assertEqual(AccessEvent.objects.count(), 1)
        self.assertFalse(Attendance.objects.filter(user=self.user).exists())

    def test_unmatched_card_is_logged_without_a_user(self):
        result = ingest_event(event('2026-10-21 08:00:00', card='9999999999', employee='999'))
        self.assertTrue(result['stored'])
        self.assertFalse(result['matched'])
        self.assertIsNone(AccessEvent.objects.get(pk=result['event_id']).user_id)
        self.assertFalse(Attendance.objects.exists())

    def test_non_person_event_is_ignored(self):
        result = ingest_event(event('2026-10-21 08:00:00', card='', employee='', minor=22))
        self.assertFalse(result['stored'])
        self.assertEqual(AccessEvent.objects.count(), 0)


@override_settings(ATTENDANCE_WEBHOOK_SECRET='testsecret', ATTENDANCE_ALLOWED_IPS='')
class WebhookTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')
        AccessCard.objects.create(user=self.user, card_no='1000000001', device_employee_no='13')

    def post(self, secret, body):
        return self.client.post(f'/attendance/hook/{secret}/', data=json.dumps(body),
                                content_type='application/json')

    def test_valid_event_is_accepted_and_stored(self):
        response = self.post('testsecret', event('2026-10-21 08:00:00'))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['ok'])
        self.assertEqual(AccessEvent.objects.filter(user=self.user).count(), 1)

    def test_wrong_secret_is_404(self):
        self.assertEqual(self.post('nope', event('2026-10-21 08:00:00')).status_code, 404)
        self.assertEqual(AccessEvent.objects.count(), 0)

    @override_settings(ATTENDANCE_ALLOWED_IPS='10.0.0.1')
    def test_disallowed_source_ip_is_404(self):
        self.assertEqual(self.post('testsecret', event('2026-10-21 08:00:00')).status_code, 404)


@override_settings(ATTENDANCE_WEBHOOK_SECRET='')
class WebhookDisabledTests(TestCase):
    def test_endpoint_is_404_without_a_configured_secret(self):
        response = self.client.post('/attendance/hook/anything/', data='{}', content_type='application/json')
        self.assertEqual(response.status_code, 404)


@override_settings(ATTENDANCE_DAY_START_HOUR=7)
class TabViewTests(TestCase):
    def setUp(self):
        self.chair = User.objects.create_user('rais', full_name='Rais Raisov', role='chair')
        self.employee = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')
        AccessCard.objects.create(user=self.employee, card_no='1000000001')
        ingest_event(event('2026-10-21 08:00:00'))
        ingest_event(event('2026-10-21 17:00:00'))

    def test_attendance_tab_shows_the_selected_day(self):
        self.client.force_login(self.chair)
        response = self.client.get('/employees/attendance/?date=2026-10-21')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'SINOV XODIM')
        self.assertContains(response, 'Kirishlar soni')

    def test_attendance_excel_export(self):
        self.client.force_login(self.chair)
        response = self.client.get('/employees/attendance/?date=2026-10-21&format=xlsx')
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

    def test_events_tab_lists_events_in_the_range(self):
        self.client.force_login(self.chair)
        response = self.client.get('/employees/events/?from=2026-10-21T00:00&to=2026-10-21T23:59')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '1000000001')

    def test_day_strip_follows_the_date_and_partial_returns_only_the_block(self):
        self.client.force_login(self.chair)
        page = self.client.get('/employees/attendance/?date=2026-10-21')
        self.assertContains(page, 'data-att-root')
        self.assertContains(page, 'data-att-day="2026-10-21" aria-current="date"')
        older = self.client.get('/employees/attendance/?date=2026-09-01')
        self.assertContains(older, 'data-att-day="2026-09-01"')  # the strip moved to the picked week
        block = self.client.get('/employees/attendance/?date=2026-10-21&partial=1')
        self.assertContains(block, 'data-att-block')
        self.assertContains(block, 'SINOV XODIM')
        self.assertNotContains(block, 'data-att-root')  # just the block, no page around it
        self.assertNotContains(block, '<html')
        future = self.client.get('/employees/attendance/?date=2999-01-01')
        self.assertNotContains(future, 'data-att-day="2999')  # clamped to today

    def test_head_sees_attendance_without_the_face_admin_tabs(self):
        head = User.objects.create_user('boshliq', full_name='Bo‘lim Boshlig‘i', role='head')
        self.client.force_login(head)
        page = self.client.get('/employees/attendance/')
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, 'href="/employees/reviews/"')
        self.assertNotContains(page, 'href="/employees/faces/"')

    def test_management_tabs_are_closed_to_a_plain_employee(self):
        self.client.force_login(self.employee)
        self.assertEqual(self.client.get('/employees/attendance/').status_code, 403)
        self.assertEqual(self.client.get('/employees/events/').status_code, 403)


@override_settings(ATTENDANCE_DAY_START_HOUR=7)
class ArrivalAlertTests(TestCase):
    def setUp(self):
        self.secretary = User.objects.create_user('kotib', full_name='Kotib Referent', role='secretary', telegram_id=1001)
        self.employee = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')
        AccessCard.objects.create(user=self.employee, card_no='1000000001')

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message')
    def test_first_entry_notifies_the_secretary_once(self, message, configured):
        with self.captureOnCommitCallbacks(execute=True):
            ingest_event(event('2026-10-21 08:00:00'))
        to_secretary = [call for call in message.call_args_list if call.args[0] == self.secretary]
        self.assertEqual(len(to_secretary), 1)
        self.assertIn('SINOV XODIM', to_secretary[0].args[1])
        self.assertTrue(AttendanceAlert.objects.filter(subject=self.employee, kind='arrival').exists())
        message.reset_mock()
        with self.captureOnCommitCallbacks(execute=True):
            ingest_event(event('2026-10-21 12:00:00'))
        self.assertEqual(message.call_count, 0)


def workday_at(hour, minute=0):
    """A fixed weekday moment, so the hourly alerts do not depend on when tests run."""
    from datetime import datetime, timedelta
    day = timezone.localdate()
    while day.weekday() >= 5:
        day -= timedelta(days=1)
    return timezone.make_aware(datetime(day.year, day.month, day.day, hour, minute))


@override_settings(ATTENDANCE_LATE_ALERT_HOUR=10, ATTENDANCE_ABSENT_HOUR=11, ATTENDANCE_WEEKEND_DAYS={5, 6})
class ScheduledAlertTests(TestCase):
    def setUp(self):
        self.secretary = User.objects.create_user('kotib', full_name='Kotib Referent', role='secretary', telegram_id=1001)
        self.came = User.objects.create_user('keldi', full_name='Keldi Kelibdi', role='employee', telegram_id=1002)
        self.absent = User.objects.create_user('yoq', full_name='Kelmadi Yoq', role='employee', telegram_id=1003)
        AccessCard.objects.create(user=self.came, card_no='111')
        AccessCard.objects.create(user=self.absent, card_no='222')
        self.day = workday_at(8).date()
        moment = workday_at(8, 30)
        Attendance.objects.create(user=self.came, date=self.day, entry_count=1,
                                  first_entry_at=moment, last_entry_at=moment)

    def run_at(self, moment):
        with patch('django.utils.timezone.localtime', return_value=timezone.localtime(moment)):
            call_command('attendance_alerts', stdout=io.StringIO())

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_late_reminder_between_ten_and_eleven(self, message, configured):
        self.run_at(workday_at(9, 30))
        self.assertEqual(message.call_count, 0)        # 09:30: too early to remind
        self.run_at(workday_at(10, 5))
        self.assertEqual([call.args[0] for call in message.call_args_list], [self.absent])
        self.assertIn('kechikyapsiz', message.call_args.args[1])

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_at_eleven_the_absentee_and_the_secretary_are_told(self, message, configured):
        self.run_at(workday_at(11, 5))
        sent = {call.args[0]: call.args[1] for call in message.call_args_list}
        self.assertIn('kelmagan deb belgilandingiz', sent[self.absent])
        self.assertNotIn('kechikyapsiz', sent[self.absent])   # no late reminder in the same breath
        self.assertIn('Kelmadi Yoq', sent[self.secretary])
        self.assertNotIn(self.came, sent)
        message.reset_mock()
        self.run_at(workday_at(12, 5))                  # idempotent: nothing new on a re-run
        self.assertEqual(message.call_count, 0)

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_a_rejected_day_counts_as_absent(self, message, configured):
        Attendance.objects.filter(user=self.came).update(status='rejected')
        self.run_at(workday_at(11, 5))
        recipients = [call.args[0] for call in message.call_args_list]
        self.assertIn(self.came, recipients)
        summary = next(call.args[1] for call in message.call_args_list if call.args[0] == self.secretary)
        self.assertIn('Keldi Kelibdi (rad etilgan)', summary)

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_weekend_is_quiet(self, message, configured):
        from datetime import timedelta
        saturday = workday_at(11, 5)
        while saturday.weekday() != 5:
            saturday += timedelta(days=1)
        self.run_at(saturday)
        self.assertEqual(message.call_count, 0)


@override_settings(ATTENDANCE_DAY_START_HOUR=7, FACE_MATCH_THRESHOLD=0.35,
                   TELEGRAM_BOT_TOKEN='', MEDIA_ROOT=tempfile.mkdtemp())
class FaceCheckTests(TestCase):
    def setUp(self):
        self.secretary = User.objects.create_user('kotib', full_name='Kotib Referent', role='secretary')
        self.employee = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')
        AccessCard.objects.create(user=self.employee, card_no='1000000001')
        FaceEnrollment.objects.create(user=self.employee, embedding=[0.1] * 512,
                                      photo=SimpleUploadedFile('enrolled.jpg', b'enrolled-bytes'))
        # The face subsystem is "up" unless a test says otherwise; no model files needed.
        available = patch('core.face.available', return_value=True)
        available.start()
        self.addCleanup(available.stop)

    def picture(self):
        return SimpleUploadedFile('capture.jpg', b'captured-bytes', content_type='image/jpeg')

    @patch('core.face.compare', return_value=0.9)
    def test_high_confidence_marks_present_without_storing_photo(self, compare):
        with self.captureOnCommitCallbacks(execute=True):
            result = ingest_event(event('2026-10-21 08:00:00'), self.picture())
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'present')
        self.assertFalse(FaceReview.objects.exists())
        self.assertFalse(bool(AccessEvent.objects.get(pk=result['event_id']).photo))

    @patch('core.attendance.send_review_alert')
    @patch('core.face.compare', return_value=0.1)
    def test_low_confidence_opens_review_and_keeps_photo(self, compare, alert):
        with self.captureOnCommitCallbacks(execute=True):
            result = ingest_event(event('2026-10-21 08:00:00'), self.picture())
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'review')
        review = FaceReview.objects.get()
        self.assertEqual(review.user, self.employee)
        self.assertTrue(bool(AccessEvent.objects.get(pk=result['event_id']).photo))
        alert.assert_called_once()

    @patch('core.telegram.message')
    @patch('core.telegram.configured', return_value=True)
    def test_employee_is_told_when_the_day_goes_to_review(self, configured, message):
        with patch('core.face.compare', return_value=None):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 08:00:00'), self.picture())
        to_employee = [call.args[1] for call in message.call_args_list if call.args[0] == self.employee]
        self.assertEqual(len(to_employee), 1)
        self.assertIn('08:00', to_employee[0])
        self.assertIn('yuzingiz topilmadi', to_employee[0])
        self.assertIn('tekshiruviga yuborildi', to_employee[0])

    @patch('core.attendance.send_review_alert')
    @patch('core.face.compare', return_value=None)
    def test_card_without_a_face_in_frame_goes_to_review(self, compare, alert):
        with self.captureOnCommitCallbacks(execute=True):
            result = ingest_event(event('2026-10-21 08:00:00'), self.picture())
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'review')
        review = FaceReview.objects.get()
        self.assertIsNone(review.confidence)
        self.assertTrue(bool(AccessEvent.objects.get(pk=result['event_id']).photo))
        alert.assert_called_once()

    @patch('core.face.compare')
    def test_face_subsystem_down_falls_back_to_card_only(self, compare):
        with patch('core.face.available', return_value=False):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 08:00:00'), self.picture())
        compare.assert_not_called()
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'present')
        self.assertFalse(FaceReview.objects.exists())

    @patch('core.telegram.message')
    @patch('core.telegram.configured', return_value=True)
    def test_review_alert_reads_well_when_no_face_was_found(self, configured, message):
        from core.attendance import send_review_alert
        with patch('core.face.compare', return_value=None):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())
        send_review_alert(FaceReview.objects.get().pk)
        self.assertIn('yuz topilmadi', message.call_args.args[1])

    @patch('core.face.compare')
    def test_no_enrollment_falls_back_to_card_only(self, compare):
        FaceEnrollment.objects.all().delete()
        with self.captureOnCommitCallbacks(execute=True):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())
        compare.assert_not_called()
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'present')
        self.assertFalse(FaceReview.objects.exists())

    @patch('core.attendance.send_not_counted')
    def test_reject_marks_rejected_and_warns_the_employee(self, warn):
        with patch('core.face.compare', return_value=0.1):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 08:00:00'), self.picture())
        review = FaceReview.objects.get()
        with self.captureOnCommitCallbacks(execute=True):
            resolve_review(self.secretary, review.pk, approve=False, reason='Kartani boshqa odam ishlatdi')
        review.refresh_from_db()
        self.assertEqual(review.state, 'rejected')
        self.assertEqual(review.reason, 'Kartani boshqa odam ishlatdi')
        self.assertEqual(review.resolved_by, self.secretary)
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'rejected')
        warn.assert_called_once()
        self.assertEqual(warn.call_args.args[2], 'Kartani boshqa odam ishlatdi')

    def test_face_tabs_and_review_badge_live_under_attendance(self):
        with patch('core.face.compare', return_value=0.1):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())
        self.client.force_login(self.secretary)
        page = self.client.get('/employees/attendance/?date=2026-10-21')
        self.assertContains(page, 'Kunlik davomat')
        self.assertContains(page, 'href="/employees/faces/"')
        self.assertContains(page, 'href="/employees/reviews/"')
        self.assertContains(page, '<span class="count-pill">1</span>')
        stats = self.client.get('/employees/')
        self.assertNotContains(stats, 'href="/employees/faces/"')  # not a top-level tab

    def reject_first_entry(self):
        with patch('core.face.compare', return_value=None):
            ingest_event(event('2026-10-21 11:09:25'), self.picture())
        resolve_review(self.secretary, FaceReview.objects.get().pk, approve=False, reason='Yuz ko‘rsatilmadi')

    @patch('core.attendance.send_reverified')
    def test_clear_face_after_a_rejection_restores_the_day_automatically(self, reverified):
        self.reject_first_entry()
        with patch('core.face.compare', return_value=0.63):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 11:38:56'), self.picture())
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'present')
        automatic = FaceReview.objects.get(event__happened_at__minute=38)
        self.assertEqual(automatic.state, 'approved')
        self.assertIsNone(automatic.resolved_by)
        self.assertIn('avtomatik', automatic.reason)
        self.assertEqual(FaceReview.objects.filter(state='rejected').count(), 1)  # the rejection stays on record
        reverified.assert_called_once()

    @patch('core.attendance.send_review_alert')
    def test_unclear_face_after_a_rejection_opens_a_new_review(self, alert):
        self.reject_first_entry()
        alert.reset_mock()
        with patch('core.face.compare', return_value=0.1):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 11:38:56'), self.picture())
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'review')
        self.assertEqual(FaceReview.objects.filter(state='pending').count(), 1)
        alert.assert_called_once()  # a second review that day still reaches the secretary

    def open_review_on_first_entry(self):
        with patch('core.face.compare', return_value=0.1):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())

    def arrival(self):
        day = Attendance.objects.get(user=self.employee)
        return day.status, timezone.localtime(day.first_entry_at).strftime('%H:%M')

    @patch('core.attendance.send_reverified')
    def test_clear_face_while_in_review_counts_but_leaves_the_earlier_review_open(self, reverified):
        self.open_review_on_first_entry()
        with patch('core.face.compare', return_value=0.7):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 09:00:00'), self.picture())
        self.assertEqual(self.arrival(), ('present', '09:00'))  # arrived at the confirmed entry
        earlier = FaceReview.objects.get(event__happened_at__hour=8)
        self.assertEqual(earlier.state, 'pending')  # still the secretary's to decide
        self.assertTrue(bool(earlier.event.photo))
        automatic = FaceReview.objects.get(event__happened_at__hour=9)
        self.assertEqual(automatic.state, 'approved')
        self.assertIsNone(automatic.resolved_by)
        reverified.assert_called_once()
        self.assertEqual(reverified.call_args.args[2], 'review')

    def test_approving_the_earlier_review_moves_the_arrival_back(self):
        self.open_review_on_first_entry()
        with patch('core.face.compare', return_value=0.7):
            ingest_event(event('2026-10-21 09:00:00'), self.picture())
        resolve_review(self.secretary, FaceReview.objects.get(event__happened_at__hour=8).pk, approve=True)
        self.assertEqual(self.arrival(), ('present', '08:00'))

    @patch('core.attendance.send_not_counted')
    def test_rejecting_the_earlier_review_keeps_a_day_already_confirmed(self, warn):
        self.open_review_on_first_entry()
        with patch('core.face.compare', return_value=0.7):
            ingest_event(event('2026-10-21 09:00:00'), self.picture())
        with self.captureOnCommitCallbacks(execute=True):
            resolve_review(self.secretary, FaceReview.objects.get(event__happened_at__hour=8).pk,
                           approve=False, reason='Boshqa odam')
        self.assertEqual(self.arrival(), ('present', '09:00'))
        warn.assert_not_called()  # the day still stands, so nothing to warn about

    @patch('core.attendance.send_review_alert')
    def test_every_unclear_entry_gets_its_own_review(self, alert):
        self.open_review_on_first_entry()
        alert.reset_mock()
        with patch('core.face.compare', return_value=0.1):
            with self.captureOnCommitCallbacks(execute=True):
                result = ingest_event(event('2026-10-21 09:00:00'), self.picture())
        self.assertEqual(FaceReview.objects.filter(state='pending').count(), 2)
        self.assertEqual(self.arrival()[0], 'review')
        self.assertTrue(bool(AccessEvent.objects.get(pk=result['event_id']).photo))
        alert.assert_called_once()
        # Approving the first one makes it the arrival, even with the second still open.
        resolve_review(self.secretary, FaceReview.objects.get(event__happened_at__hour=8).pk, approve=True)
        self.assertEqual(self.arrival(), ('present', '08:00'))

    def test_later_entries_on_a_good_day_are_not_rechecked(self):
        with patch('core.face.compare', return_value=0.9):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())
        with patch('core.face.compare') as compare:
            ingest_event(event('2026-10-21 12:00:00'), self.picture())
        compare.assert_not_called()
        self.assertEqual(Attendance.objects.get(user=self.employee).entry_count, 2)

    def test_reject_without_a_reason_is_refused(self):
        from django.core.exceptions import ValidationError
        with patch('core.face.compare', return_value=0.1):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())
        review = FaceReview.objects.get()
        with self.assertRaises(ValidationError):
            resolve_review(self.secretary, review.pk, approve=False, reason='   ')
        review.refresh_from_db()
        self.assertEqual(review.state, 'pending')

    def test_rejected_day_is_shown_as_rejected_with_its_reason(self):
        with patch('core.face.compare', return_value=0.1):
            ingest_event(event('2026-10-21 08:00:00'), self.picture())
        resolve_review(self.secretary, FaceReview.objects.get().pk, approve=False, reason='Yuz mos emas')
        self.client.force_login(self.secretary)
        page = self.client.get('/employees/attendance/?date=2026-10-21')
        self.assertContains(page, 'Rad etildi')
        self.assertContains(page, 'Yuz mos emas')
        self.assertContains(page, 'Keldi: <strong>0</strong>')
        self.assertContains(page, 'Rad etildi: <strong>1</strong>')
        history = self.client.get('/employees/reviews/')
        self.assertContains(history, 'Yuz mos emas')
        export = self.client.get('/employees/attendance/?date=2026-10-21&format=xlsx')
        self.assertEqual(export.status_code, 200)

    def test_approve_marks_present(self):
        with patch('core.face.compare', return_value=0.1):
            with self.captureOnCommitCallbacks(execute=True):
                ingest_event(event('2026-10-21 08:00:00'), self.picture())
        review = FaceReview.objects.get()
        resolve_review(self.secretary, review.pk, approve=True)
        self.assertEqual(Attendance.objects.get(user=self.employee).status, 'present')


@override_settings(TELEGRAM_BOT_TOKEN='', MEDIA_ROOT=tempfile.mkdtemp())
class EnrollmentTests(TestCase):
    def setUp(self):
        self.chair = User.objects.create_user('rais', full_name='Rais Raisov', role='chair')
        self.employee = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')

    @patch('core.face.embed', return_value=[0.2] * 512)
    def test_enroll_stores_embedding_and_photo(self, embed):
        from core.attendance import enroll_face
        enroll_face(self.employee, SimpleUploadedFile('f.jpg', b'photo-bytes'))
        enrollment = FaceEnrollment.objects.get(user=self.employee)
        self.assertEqual(len(enrollment.embedding), 512)
        self.assertTrue(bool(enrollment.photo))

    def test_review_page_access(self):
        self.client.force_login(self.chair)
        self.assertEqual(self.client.get('/employees/reviews/').status_code, 200)
        self.assertEqual(self.client.get('/employees/faces/').status_code, 200)
        self.client.force_login(self.employee)
        self.assertEqual(self.client.get('/employees/reviews/').status_code, 403)


@override_settings(ATTENDANCE_DAY_START_HOUR=7, ATTENDANCE_LATE_HOUR=9, TELEGRAM_BOT_TOKEN='')
class LatenessAndExcuseTests(TestCase):
    def setUp(self):
        self.secretary = User.objects.create_user('kotib', full_name='Kotib Referent', role='secretary')
        self.employee = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee')
        self.other = User.objects.create_user('boshqa', full_name='Boshqa Xodim', role='employee')
        AccessCard.objects.create(user=self.employee, card_no='1000000001')
        self.client.force_login(self.secretary)

    def page(self, day='2026-10-21'):
        return self.client.get(f'/employees/attendance/?date={day}')

    def excuse(self, **data):
        body = {'employee': self.other.pk, 'date': '2026-10-21', 'kind': 'sick', **data}
        return self.client.post('/employees/attendance/excuse/', body)

    def test_arrival_after_nine_is_late_with_its_delay(self):
        ingest_event(event('2026-10-21 10:25:00'))
        from core.attendance import late_minutes
        self.assertEqual(late_minutes(Attendance.objects.get(user=self.employee)), 85)
        page = self.page()
        self.assertContains(page, '+1 soat 25 daq')
        self.assertContains(page, 'Kechikdi: <strong>1</strong>')

    def test_arrival_before_nine_is_not_late(self):
        ingest_event(event('2026-10-21 08:59:00'))
        from core.attendance import late_minutes
        self.assertIsNone(late_minutes(Attendance.objects.get(user=self.employee)))
        self.assertContains(self.page(), 'Kechikdi: <strong>0</strong>')

    def test_a_day_in_review_has_no_lateness_yet(self):
        ingest_event(event('2026-10-21 10:00:00'))
        Attendance.objects.filter(user=self.employee).update(status='review')
        from core.attendance import late_minutes
        self.assertIsNone(late_minutes(Attendance.objects.get(user=self.employee)))

    def test_excuse_for_several_days_shows_on_the_attendance_tab(self):
        response = self.excuse(until='2026-10-22', note='Shamollagan')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.other.attendance_excuses.count(), 2)
        excuse = self.other.attendance_excuses.first()
        self.assertEqual(excuse.created_by, self.secretary)
        page = self.page()
        self.assertContains(page, 'Kasal: Shamollagan')
        self.assertContains(page, 'Uzrli: <strong>1</strong>')

    def test_coming_later_keeps_the_expected_time(self):
        self.excuse(kind='late', expected_at='11:00', note='11 dan keyin kelaman dedi')
        excuse = self.other.attendance_excuses.get()
        self.assertEqual(excuse.label, 'Kechroq keladi (11:00 da): 11 dan keyin kelaman dedi')

    def test_other_reason_needs_a_note_and_dates_must_run_forward(self):
        self.excuse(kind='other')
        self.excuse(until='2026-10-20')
        self.assertFalse(AttendanceExcuse.objects.exists())

    def test_excuse_can_be_cancelled(self):
        self.excuse()
        excuse = self.other.attendance_excuses.get()
        self.client.post(f'/employees/attendance/excuse/{excuse.pk}/delete/')
        self.assertFalse(AttendanceExcuse.objects.exists())

    def test_excuses_stay_within_the_users_own_scope(self):
        head = User.objects.create_user('boshliq', full_name='Bo‘lim Boshlig‘i', role='head')
        self.client.force_login(head)
        self.assertEqual(self.excuse().status_code, 404)  # not one of his people
        self.client.force_login(self.employee)
        self.assertEqual(self.excuse().status_code, 403)
        self.assertFalse(AttendanceExcuse.objects.exists())

    def test_excel_carries_lateness_and_excuse(self):
        ingest_event(event('2026-10-21 10:25:00'))
        self.excuse(note='Shamollagan')
        response = self.client.get('/employees/attendance/?date=2026-10-21&format=xlsx')
        self.assertEqual(response.status_code, 200)
        from io import BytesIO
        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(response.content)).active
        cells = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
        self.assertIn('+1 soat 25 daq', cells)
        self.assertIn('Kasal: Shamollagan', cells)


@override_settings(ATTENDANCE_LATE_HOUR=0, ATTENDANCE_ABSENT_HOUR=0, ATTENDANCE_DAY_START_HOUR=7)
class ExcusedAlertTests(TestCase):
    def setUp(self):
        self.secretary = User.objects.create_user('kotib', full_name='Kotib Referent', role='secretary', telegram_id=1001)
        self.sick = User.objects.create_user('kasal', full_name='Kasal Xodim', role='employee', telegram_id=1002)
        self.missing = User.objects.create_user('yoq', full_name='Kelmadi Yoq', role='employee', telegram_id=1003)
        AccessCard.objects.create(user=self.sick, card_no='111')
        AccessCard.objects.create(user=self.missing, card_no='222')
        AttendanceExcuse.objects.create(user=self.sick, date=timezone.localdate(), kind='sick', note='Shamollagan')

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_excused_employee_gets_no_late_warning_and_is_listed_apart(self, message, configured):
        call_command('attendance_alerts')
        recipients = [call.args[0] for call in message.call_args_list]
        self.assertNotIn(self.sick, recipients)
        self.assertIn(self.missing, recipients)
        summary = next(call.args[1] for call in message.call_args_list if call.args[0] == self.secretary)
        self.assertIn('kelmaganlar (1)', summary)
        self.assertIn('Uzrli (1)', summary)
        self.assertIn('Kasal Xodim — Kasal: Shamollagan', summary)


@override_settings(ATTENDANCE_DAY_START_HOUR=7, ATTENDANCE_LATE_HOUR=9, FACE_MATCH_THRESHOLD=0.35,
                   ATTENDANCE_EXEMPT_ROLES=['chair'])
class CheckedInReceiptTests(TestCase):
    def setUp(self):
        self.employee = User.objects.create_user('anas', full_name='SINOV XODIM', role='employee', telegram_id=3001)
        AccessCard.objects.create(user=self.employee, card_no='1000000001')
        FaceEnrollment.objects.create(user=self.employee, embedding=[0.1] * 512,
                                      photo=SimpleUploadedFile('enrolled.jpg', b'enrolled-bytes'))
        available = patch('core.face.available', return_value=True)
        available.start()
        self.addCleanup(available.stop)

    def picture(self):
        return SimpleUploadedFile('capture.jpg', b'captured-bytes', content_type='image/jpeg')

    def employee_messages(self, message):
        return [call.args[1] for call in message.call_args_list if call.args[0] == self.employee]

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    @patch('core.face.compare', return_value=0.9)
    def test_recognised_arrival_sends_one_receipt_with_lateness(self, compare, message, configured):
        with self.captureOnCommitCallbacks(execute=True):
            ingest_event(event('2026-10-21 09:25:00'), self.picture())
        texts = self.employee_messages(message)
        self.assertEqual(len(texts), 1)
        self.assertIn('Yo‘qlamangiz qabul qilindi', texts[0])
        self.assertIn('09:25', texts[0])
        self.assertIn('25 daq', texts[0])
        with self.captureOnCommitCallbacks(execute=True):
            ingest_event(event('2026-10-21 13:00:00'), self.picture())
        self.assertEqual(len(self.employee_messages(message)), 1)   # once a day

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    @patch('core.face.compare', return_value=0.1)
    def test_secretary_approval_sends_the_receipt(self, compare, message, configured):
        with self.captureOnCommitCallbacks(execute=True):
            ingest_event(event('2026-10-21 08:40:00'), self.picture())
        self.assertFalse(any('qabul qilindi' in t for t in self.employee_messages(message)))
        secretary = User.objects.create_user('kotib', full_name='Kotib', role='secretary')
        with self.captureOnCommitCallbacks(execute=True):
            resolve_review(secretary, FaceReview.objects.get().pk, True)
        self.assertTrue(any('qabul qilindi' in t for t in self.employee_messages(message)))
