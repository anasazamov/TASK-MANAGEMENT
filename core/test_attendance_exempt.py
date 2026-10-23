import io
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .attendance import ingest_event
from .models import AccessCard, AccessEvent, Attendance, User


def swipe(card, time='2026-10-21 08:10:00'):
    return {'major': 5, 'minor': 1, 'time': time, 'cardNo': card, 'employeeNoString': '', 'name': 'X', 'doorNo': 1}


@override_settings(ATTENDANCE_DAY_START_HOUR=7, ATTENDANCE_DOOR_MAJOR=5, ATTENDANCE_EXEMPT_ROLES=['chair'])
class ChairIsExemptTests(TestCase):
    def setUp(self):
        self.chair = User.objects.create_user('rais', full_name='Sinov Rais', role='chair', telegram_id=2001, password='x')
        self.secretary = User.objects.create_user('kotib', full_name='Sinov Kotib', role='secretary',
                                                  telegram_id=2002, password='x')
        self.worker = User.objects.create_user('xodim', full_name='Sinov Xodim', role='employee', telegram_id=2003)
        AccessCard.objects.create(user=self.chair, card_no='1000000001')
        AccessCard.objects.create(user=self.worker, card_no='1000000002')

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_chair_entry_is_logged_but_not_attendance(self, message, configured):
        with self.captureOnCommitCallbacks(execute=True):
            result = ingest_event(swipe('1000000001'))
        self.assertEqual(result['reason'], 'exempt')
        self.assertEqual(AccessEvent.objects.get().user, self.chair)          # the door log keeps it
        self.assertFalse(Attendance.objects.exists())
        message.assert_not_called()                                           # no "arrived" to the secretary

    def test_chair_is_not_on_the_daily_list(self):
        self.client.force_login(self.secretary)
        page = self.client.get(reverse('attendance_daily'))
        names = [row['user'] for row in page.context['rows']]
        self.assertIn(self.worker, names)
        self.assertNotIn(self.chair, names)

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    def test_chair_gets_no_late_warning_and_is_not_absent(self, message, configured):
        from .test_attendance import workday_at
        with patch('django.utils.timezone.localtime', return_value=timezone.localtime(workday_at(12))):
            call_command('attendance_alerts', stdout=io.StringIO())
        recipients = [call.args[0] for call in message.call_args_list]
        self.assertNotIn(self.chair, recipients)
        summary = next(call.args[1] for call in message.call_args_list if call.args[0] == self.secretary)
        self.assertIn('Sinov Xodim', summary)
        self.assertNotIn('Sinov Rais', summary)

    def test_chair_has_no_attendance_page_in_the_menu(self):
        self.client.force_login(self.chair)
        self.assertNotContains(self.client.get(reverse('dashboard')), reverse('my_attendance'))
        self.assertContains(self.client.get(reverse('my_attendance')), 'davomatingiz yuritilmaydi')


class MarkedExportTests(TestCase):
    def setUp(self):
        self.secretary = User.objects.create_user('kotib', full_name='Sinov Kotib', role='secretary', password='x')
        self.worker = User.objects.create_user('xodim', full_name='Sinov Xodim', role='employee')
        self.client.force_login(self.secretary)

    def sheet(self, **params):
        from openpyxl import load_workbook
        response = self.client.get(reverse('attendance_daily'), {'format': 'xlsx', **params})
        return response, load_workbook(io.BytesIO(response.content)).active

    def row_of(self, sheet, name):
        return next(r for r in range(5, sheet.max_row + 1) if sheet.cell(row=r, column=2).value == name)

    def test_absentees_are_painted_and_counted(self):
        response, sheet = self.sheet(marked='1')
        self.assertIn('belgilangan', response['Content-Disposition'])
        row = self.row_of(sheet, 'Sinov Xodim')
        self.assertEqual(sheet.cell(row=row, column=8).value, 'Kelmadi')
        self.assertEqual(sheet.cell(row=row, column=2).fill.start_color.rgb[-6:], 'FFC7CE')
        legend = [sheet.cell(row=r, column=2).value for r in range(row + 1, sheet.max_row + 1)]
        self.assertIn('Kelmadi: 2', legend)          # the worker and the secretary herself

    def test_plain_export_has_no_colour(self):
        _, sheet = self.sheet()
        row = self.row_of(sheet, 'Sinov Xodim')
        self.assertNotEqual(sheet.cell(row=row, column=2).fill.fill_type, 'solid')
