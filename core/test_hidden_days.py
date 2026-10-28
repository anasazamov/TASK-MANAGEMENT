import io
from datetime import date
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from .attendance import ingest_event
from .models import AccessCard, Attendance, AttendanceAlert, FaceEnrollment, FaceReview, User

SATURDAY, SUNDAY = '2026-10-24', '2026-10-25'


def swipe(time):
    return {'major': 5, 'minor': 1, 'time': time, 'cardNo': '1000000001', 'employeeNoString': '', 'name': 'X',
            'doorNo': 1}


@override_settings(ATTENDANCE_DAY_START_HOUR=7, ATTENDANCE_LATE_HOUR=9, ATTENDANCE_DOOR_MAJOR=5,
                   ATTENDANCE_HIDDEN_DAYS={5, 6}, ATTENDANCE_EXEMPT_ROLES=['chair'], TELEGRAM_BOT_TOKEN='')
class WeekendTimesHiddenTests(TestCase):
    def setUp(self):
        self.employee = User.objects.create_user('xodim', full_name='Sinov Xodim', role='employee',
                                                 telegram_id=5001, password='x')
        self.secretary = User.objects.create_user('kotib', full_name='Sinov Kotib', role='secretary',
                                                  telegram_id=5002, password='x')
        self.root = User.objects.create_superuser('admin', password='Testing-123!')
        AccessCard.objects.create(user=self.employee, card_no='1000000001')
        FaceEnrollment.objects.create(user=self.employee, embedding=[0.1] * 512,
                                      photo=SimpleUploadedFile('f.jpg', b'f'))

    @patch('core.telegram.configured', return_value=True)
    @patch('core.telegram.message', return_value=True)
    @patch('core.face.available', return_value=True)
    @patch('core.face.best_match', return_value=0.1)
    def test_weekend_is_recorded_quietly(self, match, available, message, configured):
        picture = SimpleUploadedFile('c.jpg', b'c', content_type='image/jpeg')
        with self.captureOnCommitCallbacks(execute=True):
            result = ingest_event(swipe(f'{SUNDAY} 09:47:00'), picture)
        self.assertEqual(result['reason'], 'hidden_day')
        self.assertEqual(Attendance.objects.get(date=SUNDAY).status, 'present')
        match.assert_not_called()                     # no face check, so no review shows the time
        self.assertFalse(FaceReview.objects.exists())
        message.assert_not_called()
        self.assertFalse(AttendanceAlert.objects.exists())

    def test_secretary_sees_who_came_on_saturday_but_not_when(self):
        ingest_event(swipe(f'{SATURDAY} 09:47:00'))
        self.client.force_login(self.secretary)
        page = self.client.get(reverse('attendance_daily') + f'?date={SATURDAY}')
        self.assertEqual(str(page.context['day']), SATURDAY)
        self.assertIn(SATURDAY, [str(chip['date']) for chip in page.context['chips']])
        row = next(row for row in page.context['rows'] if row['user'] == self.employee)
        self.assertIsNotNone(row['attendance'])
        self.assertIsNone(row['late'])                       # no lateness on a hidden day
        self.assertContains(page, 'Keldi')
        self.assertNotContains(page, '09:47')
        events = self.client.get(reverse('access_events') + f'?from={SATURDAY}T00:00&to={SATURDAY}T23:59')
        self.assertEqual(len(events.context['events']), 0)   # each door event is a time

    def test_export_leaves_the_times_out(self):
        from openpyxl import load_workbook
        ingest_event(swipe(f'{SATURDAY} 09:47:00'))
        self.client.force_login(self.secretary)
        response = self.client.get(reverse('attendance_daily'), {'date': SATURDAY, 'format': 'xlsx'})
        sheet = load_workbook(io.BytesIO(response.content)).active
        row = next(r for r in range(5, sheet.max_row + 1) if sheet.cell(row=r, column=2).value == 'Sinov Xodim')
        self.assertEqual([sheet.cell(row=row, column=c).value for c in (5, 6, 7, 8)], ['—', '—', '—', 'Keldi'])

    def test_employee_month_shows_the_day_without_times(self):
        from .views import my_month_rows
        ingest_event(swipe(f'{SATURDAY} 09:47:00'))
        rows = {row['date']: row for row in my_month_rows(self.employee, date(2026, 10, 1), date(2026, 10, 25),
                                                          date(2026, 10, 20))}
        saturday = rows[date(2026, 10, 24)]
        self.assertEqual((saturday['state'], saturday['hide_times']), ('present', True))

    def test_superuser_reads_the_times_in_the_admin(self):
        ingest_event(swipe(f'{SATURDAY} 09:47:00'))
        self.client.force_login(self.root)
        self.assertContains(self.client.get(reverse('admin:core_attendance_changelist')), '09:47')
