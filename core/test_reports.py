import io
from datetime import timedelta

from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from . import reports
from .models import Department, Task, User
from .services import task_action


@override_settings(STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
                             'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class EmployeeReportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        dep, other_dep = Department.objects.create(name='Qurilish'), Department.objects.create(name='IT')
        cls.chair = User.objects.create_user('rais', full_name='Abilov Feruz', role='chair')
        cls.head = User.objects.create_user('boshliq', full_name='Tuyliyev Asliddin', role='head', department=dep)
        cls.employee = User.objects.create_user('xodim', full_name='Xalimov Farrux', department=dep)
        cls.other = User.objects.create_user('other', full_name='Maxfiy Xodim', department=other_dep)
        now = timezone.now()
        make = lambda **kw: Task.objects.create(issuer=cls.head, assignee=cls.employee, title='Ish', **kw)
        cls.future = make(due_at=now+timedelta(days=5))
        cls.late = make(due_at=now-timedelta(days=1), seen_at=now)
        cls.undated = make(seen_at=now)
        cls.submitted = make(status='submitted', seen_at=now)
        cls.accepted = make(status='accepted', seen_at=now)
        cls.old = make(created_at=now-timedelta(days=21), seen_at=now)
        Task.objects.create(issuer=cls.chair, assignee=cls.other, title='Boshqa bo‘lim')

    def row(self, user, period=None, person=None):
        return next(r for r in reports.employee_stats(user, period) if r['assignee_id'] == (person or self.employee).pk)

    def test_columns_match_task_states(self):
        row = self.row(self.chair)
        self.assertEqual((row['total'], row['in_progress'], row['overdue'], row['unseen']), (6, 4, 1, 1))

    def test_head_sees_only_visible_employees_and_tasks(self):
        self.assertEqual([r['assignee_id'] for r in reports.employee_stats(self.head)], [self.employee.pk])
        self.assertEqual({r['assignee_id'] for r in reports.employee_stats(self.chair)},
                         {self.employee.pk, self.other.pk, self.head.pk})
        self.client.force_login(self.employee)
        self.assertEqual(self.client.get('/employees/').status_code, 403)

    def test_people_without_tasks_are_listed_with_zeros(self):
        idle = reports.employee_stats(self.chair)[-1]
        self.assertEqual((idle['assignee_id'], idle['total'], idle['unseen']), (self.head.pk, 0, 0))
        office = User.objects.create_user('devon', full_name='Isomiddinov Islom', role='office')
        rows = reports.employee_stats(office)
        self.assertEqual({r['assignee_id'] for r in rows}, {self.head.pk})  # The office assigns to heads.
        self.assertEqual(rows[0]['total'], 0)

    def test_week_filter_counts_tasks_created_that_week(self):
        period = reports.parse_week(reports.current_week())
        self.assertEqual(self.row(self.chair, period)['total'], 5)
        self.assertIsNone(reports.parse_week('bad'))
        self.client.force_login(self.head)
        page = self.client.get('/employees/', {'week': reports.current_week()})
        self.assertContains(page, self.employee.full_name)
        self.assertEqual(page.context['rows'][0]['total'], 5)

    def test_assignee_opening_or_acting_marks_seen_without_touching_updated_at(self):
        before = self.future.updated_at
        self.client.force_login(self.head)
        self.client.get(self.future.get_absolute_url())
        self.future.refresh_from_db()
        self.assertIsNone(self.future.seen_at)
        self.client.force_login(self.employee)
        self.client.get(self.future.get_absolute_url())
        self.future.refresh_from_db()
        self.assertIsNotNone(self.future.seen_at)
        self.assertEqual(self.future.updated_at, before)
        fresh = Task.objects.create(issuer=self.head, assignee=self.employee, title='Yangi')
        task_action(self.employee, fresh.pk, 'comment', 'Tanishdim')
        fresh.refresh_from_db()
        self.assertIsNotNone(fresh.seen_at)

    def test_excel_download_has_sample_layout(self):
        self.client.force_login(self.chair)
        response = self.client.get('/employees/', {'format': 'xlsx'})
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment; filename="xodimlar-statistikasi-barcha-', response['Content-Disposition'])
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[5]], [title for title, _ in reports.COLUMNS])
        self.assertEqual([c.value for c in sheet[6]], [1, 'Xalimov Farrux', 6, 4, 1, 1])
        self.assertEqual(sheet['C2'].value, 'Davr: Barcha vaqt')
        signature = 5 + len(reports.employee_stats(self.chair)) + 3
        self.assertEqual(sheet.cell(row=signature, column=3).value, self.chair.short_name)
        weekly = self.client.get('/employees/', {'format': 'xlsx', 'week': reports.current_week()})
        self.assertIn(reports.current_week(), weekly['Content-Disposition'])
        self.assertEqual(load_workbook(io.BytesIO(weekly.content)).active['C6'].value, 5)
